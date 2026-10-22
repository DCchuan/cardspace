import io
import re
from pathlib import Path
from fastapi import HTTPException
from ..db import connect, uid, now, log
from .cards import snapshot, require


def parse_document(filename,blob):
    suffix=Path(filename).suffix.lower()
    if suffix in ('.txt','.md','.markdown'):
        try: text=blob.decode('utf-8-sig')
        except UnicodeDecodeError: raise HTTPException(422,'文本文件需要 UTF-8 编码。')
        return [('全文',text)]
    if suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(blob))
        if reader.is_encrypted:raise HTTPException(422,'请先解密 PDF 再导入。')
        return [(f'第 {i+1} 页',p.extract_text() or '') for i,p in enumerate(reader.pages)]
    if suffix=='.docx':
        import zipfile
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            if sum(x.file_size for x in z.infolist())>100*1024*1024:raise HTTPException(422,'DOCX 解压内容超过限制。')
        from docx import Document
        doc=Document(io.BytesIO(blob))
        return [('文档正文','\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows))]
    raise HTTPException(422,'支持 Markdown、TXT、PDF 和 DOCX。')

def import_document(filename,blob,actor='human'):
    pages=parse_document(filename,blob)
    if not any(t.strip() for _,t in pages):raise HTTPException(422,'未提取到文本。扫描版 PDF 需要先进行 OCR。')
    with connect() as db:
        source_id,sid,duplicate=snapshot(db,'document',blob,filename)
        existing=db.execute('SELECT id FROM knowledge_documents WHERE snapshot_id=?',(sid,)).fetchone()
        if duplicate and existing:return {'document_id':existing['id'],'duplicate':True}
        did=uid('doc')
        db.execute('INSERT INTO knowledge_documents VALUES(?,?,?,?,?,?,?)',(did,source_id,sid,Path(filename).stem,filename,1,now()))
        count=0
        for locator,text in pages:
            for start in range(0,len(text),1000):
                chunk=text[start:start+1200].strip()
                if not chunk:continue
                db.execute('INSERT INTO knowledge_chunks VALUES(?,?,?,?,?)',(uid('chunk'),did,chunk,f'{locator} · 字符 {start+1}–{start+len(chunk)}',count));count+=1
        log(db,actor,'document.import',did,{'chunks':count})
        return {'document_id':did,'chunks':count,'duplicate':False}
